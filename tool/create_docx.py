import sys
import subprocess

def install_and_import(package):
    try:
        __import__(package)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

install_and_import('docx')
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
title = doc.add_heading('Tài liệu Đặc tả Yêu cầu Hệ thống (SRS)', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

p = doc.add_paragraph()
p.add_run('Dự án: ').bold = True
p.add_run('Nền tảng Thương mại Điện tử và Quản trị Nexera\n')
p.add_run('Phiên bản: ').bold = True
p.add_run('1.0')
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# 1. Giới thiệu Tổng quan
doc.add_heading('1. Giới thiệu Tổng quan (Overview)', level=1)
doc.add_paragraph('Hệ thống là một nền tảng website thương mại điện tử kết hợp quản lý quan hệ khách hàng (CRM). Mục tiêu của nền tảng là cung cấp kênh truyền thông chính thức, giới thiệu các giải pháp công nghệ, năng lượng xanh, đồng thời cho phép khách hàng trải nghiệm mua sắm trực tuyến (thiết bị, vật tư, gói lắp đặt) và thanh toán tự động.')
doc.add_paragraph('Hệ thống được chia làm hai phân hệ chính:')
doc.add_paragraph('Phân hệ Khách hàng (Customer Storefront): Giao diện công khai dành cho khách hàng truy cập, tìm hiểu thông tin và mua sắm.', style='List Bullet')
doc.add_paragraph('Phân hệ Quản trị (Admin Dashboard & CRM): Khu vực nội bộ được bảo mật, dành riêng cho ban quản lý và nhân viên để vận hành hệ thống.', style='List Bullet')

# 2. Đặc tả Phân hệ Khách hàng
doc.add_heading('2. Đặc tả Phân hệ Khách hàng (Customer Storefront)', level=1)

doc.add_heading('2.1. Module Giới thiệu & Định vị thương hiệu', level=2)
doc.add_paragraph('Trang chủ: Hiển thị tổng quan về các sản phẩm chủ lực, dự án tiêu biểu và tin tức mới nhất.', style='List Bullet')
doc.add_paragraph('Về chúng tôi: Trình bày tầm nhìn, sứ mệnh, hồ sơ năng lực của công ty.', style='List Bullet')
doc.add_paragraph('Chứng nhận & Đối tác: Hiển thị các giải thưởng, giấy chứng nhận chất lượng và danh sách đối tác chiến lược.', style='List Bullet')

doc.add_heading('2.2. Module Sản phẩm & Bán hàng (E-commerce)', level=2)
doc.add_paragraph('Danh mục Sản phẩm: Phân loại rõ ràng (Thiết bị & Vật tư, Các gói lắp đặt trọn gói,...). Khách hàng có thể tìm kiếm, lọc sản phẩm theo mức giá, công suất, nhu cầu.', style='List Bullet')
doc.add_paragraph('Chi tiết Sản phẩm: Hiển thị hình ảnh, thông số kỹ thuật chi tiết, giá bán, chính sách bảo hành.', style='List Bullet')
doc.add_paragraph('Giỏ hàng trực tuyến: Hỗ trợ khách hàng thêm/sửa/xóa sản phẩm, tự động tính tổng tiền. Giữ nguyên giỏ hàng ngay cả khi khách hàng thoát trang.', style='List Bullet')
doc.add_paragraph('Thanh toán (Checkout): Khách hàng nhập thông tin giao hàng/lắp đặt. Tích hợp cổng thanh toán trực tuyến (thẻ ngân hàng, mã QR, ví điện tử). Tự động hiển thị màn hình xác nhận giao dịch thành công.', style='List Bullet')

doc.add_heading('2.3. Module Truyền thông & Dự án', level=2)
doc.add_paragraph('Dự án tiêu biểu: Trưng bày các dự án đã thực hiện, phân loại theo nhóm (Dự án công nghiệp, Dự án dân dụng). Mỗi dự án có hình ảnh thực tế và thông tin công suất.', style='List Bullet')
doc.add_paragraph('Tin tức & Kiến thức: Chuyên mục blog cập nhật tin tức nội bộ công ty, xu hướng thị trường năng lượng, bài báo truyền thông và mẹo kinh nghiệm hữu ích.', style='List Bullet')

doc.add_heading('2.4. Module Liên hệ & Thu thập Khách hàng tiềm năng (Leads)', level=2)
doc.add_paragraph('Cổng liên hệ: Cung cấp thông tin địa chỉ, hotline, email.', style='List Bullet')
doc.add_paragraph('Biểu mẫu đăng ký tư vấn: Khách hàng để lại thông tin (Tên, SĐT, Nhu cầu lắp đặt, Tiền điện hàng tháng). Hệ thống sẽ tự động ghi nhận và chuyển thông tin này về phân hệ CRM cho bộ phận Sale.', style='List Bullet')

# 3. Đặc tả Phân hệ Quản trị
doc.add_heading('3. Đặc tả Phân hệ Quản trị (Admin Dashboard & CRM)', level=1)

doc.add_heading('3.1. Quản lý Sản phẩm & Nội dung (CMS)', level=2)
doc.add_paragraph('Quản lý Sản phẩm: Thêm mới, chỉnh sửa, xóa bỏ thông tin sản phẩm, cập nhật giá bán, hình ảnh và trạng thái còn/hết hàng.', style='List Bullet')
doc.add_paragraph('Quản lý Bài viết & Dự án: Công cụ soạn thảo văn bản trực quan giúp nhân viên dễ dàng đăng tải tin tức mới, cập nhật dự án tiêu biểu mà không cần can thiệp kỹ thuật.', style='List Bullet')

doc.add_heading('3.2. Quản lý Đơn hàng & Giao dịch', level=2)
doc.add_paragraph('Theo dõi Đơn hàng: Xem danh sách toàn bộ đơn hàng phát sinh từ website.', style='List Bullet')
doc.add_paragraph('Trạng thái Đơn hàng: Nhân viên có thể cập nhật trạng thái đơn (Chờ xử lý, Đã thanh toán, Đang vận chuyển, Hoàn tất).', style='List Bullet')
doc.add_paragraph('Lịch sử Giao dịch: Đối soát các khoản thanh toán từ cổng thanh toán tự động.', style='List Bullet')

doc.add_heading('3.3. Quản trị Quan hệ Khách hàng (CRM)', level=2)
doc.add_paragraph('Hồ sơ Khách hàng: Quản lý danh sách khách hàng đã mua hàng hoặc để lại thông tin tư vấn.', style='List Bullet')
doc.add_paragraph('Tiếp nhận Leads: Biểu mẫu tư vấn từ website sẽ tự động tạo thành một "Lead" (Cơ hội bán hàng) mới trong CRM để nhân viên Sale chủ động liên hệ.', style='List Bullet')
doc.add_paragraph('Lịch sử Khách hàng: Xem lại lịch sử các đơn hàng khách đã mua, số tiền đã chi tiêu.', style='List Bullet')

doc.add_heading('3.4. Thống kê & Báo cáo (Analytics)', level=2)
doc.add_paragraph('Bảng điều khiển (Dashboard): Biểu đồ trực quan thống kê tổng doanh thu theo ngày/tháng/năm.', style='List Bullet')
doc.add_paragraph('Báo cáo kinh doanh: Thống kê số lượng đơn hàng, top các sản phẩm bán chạy nhất để ra quyết định nhập hàng.', style='List Bullet')

# 4. Yêu cầu Phi chức năng
doc.add_heading('4. Yêu cầu Phi chức năng (Yêu cầu chung)', level=1)
doc.add_paragraph('Giao diện đa nền tảng (Responsive): Website phải hiển thị tốt và thao tác dễ dàng trên mọi thiết bị (Điện thoại di động, Máy tính bảng, PC).', style='List Number')
doc.add_paragraph('Trải nghiệm người dùng (UX/UI): Thiết kế phải mang phong cách hiện đại, trực quan, dễ sử dụng, thể hiện được sự chuyên nghiệp của định vị "Công nghệ thông minh & Năng lượng xanh".', style='List Number')
doc.add_paragraph('Bảo mật dữ liệu: Thông tin khách hàng và lịch sử giao dịch phải được mã hóa và bảo mật nghiêm ngặt. Khu vực quản trị phải được bảo vệ bằng tài khoản và mật khẩu phân quyền.', style='List Number')
doc.add_paragraph('Tốc độ phản hồi: Thời gian tải trang phải được tối ưu ở mức nhanh nhất, đảm bảo khách hàng không phải chờ đợi lâu khi duyệt sản phẩm, đồng thời đáp ứng tốt tiêu chuẩn chấm điểm SEO của Google.', style='List Number')

doc.save('Nexera_SRS.docx')
print("Tạo file Nexera_SRS.docx thành công!")
